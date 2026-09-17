from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


STATUS_LABELS = {
    "WAITING": "候补中",
    "SELECTED": "容量内",
    "FINAL": "已确认",
    "REJECTED": "未录取",
}

PERIOD_STATUS_LABELS = {
    "WAITING": "未开始",
    "ACTIVE": "进行中",
    "CLOSED": "已结束",
}

TITLE_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FILL = PatternFill("solid", fgColor="D9EDE9")
STRIPE_FILL = PatternFill("solid", fgColor="F3F8F7")
WARNING_FILL = PatternFill("solid", fgColor="FDE2E2")
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _style_title(sheet, title: str, last_column: int):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = sheet.cell(row=1, column=1, value=title)
    cell.data_type = "s"
    cell.font = Font(color="FFFFFF", bold=True, size=16)
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30


def _style_metadata(sheet, text: str, last_column: int):
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    cell = sheet.cell(row=2, column=1, value=text)
    cell.data_type = "s"
    cell.font = Font(color="475569", size=10)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 22


def _write_table(sheet, headers: list[str], rows: list[list], widths: list[int]):
    header_row = 4
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=value)
        cell.font = Font(bold=True, color="134E4A")
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, values in enumerate(rows, start=header_row + 1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # Never interpret student/course data as Excel formulas.
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="center" if column != len(headers) else "left",
                vertical="center",
            )
            if row_index % 2 == 0:
                cell.fill = STRIPE_FILL
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width

    sheet.freeze_panes = "A5"
    if rows:
        sheet.auto_filter.ref = f"A4:{chr(64 + len(headers))}{header_row + len(rows)}"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _period_metadata(period):
    start_time = period.start_time.strftime("%Y-%m-%d %H:%M:%S")
    end_time = period.end_time.strftime("%Y-%m-%d %H:%M:%S")
    status = PERIOD_STATUS_LABELS.get(period.status, period.status)
    return f"阶段时间：{start_time} 至 {end_time}    阶段状态：{status}"


def _ranking_rows(course, ranking):
    rows = []
    for rank, item in enumerate(ranking, start=1):
        status = item.Selection.status
        note = ""
        display_rank = "-" if status == "REJECTED" else rank
        if status == "REJECTED":
            note = "未录取（不占用后续选课资格）"
        elif status == "WAITING":
            note = "候补"
        elif rank > course.capacity:
            note = "超出容量"
        rows.append([
            course.name,
            course.instructor or "",
            course.schedule or "",
            course.location or "",
            display_rank,
            item.Student.student_no,
            item.Student.name,
            item.Student.weight,
            STATUS_LABELS.get(status, status),
            item.Selection.selected_time,
            note,
        ])
    return rows


def build_period_selection_workbook(period, courses, rankings_by_course):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "课程汇总"
    _style_title(summary, f"{period.name} · 选课情况汇总", 12)
    _style_metadata(summary, _period_metadata(period), 12)

    summary_rows = []
    detail_rows = []
    for index, course in enumerate(courses, start=1):
        ranking = rankings_by_course.get(course.id, [])
        admitted_count = sum(
            1 for item in ranking if item.Selection.status in ("SELECTED", "FINAL")
        )
        waiting_count = sum(
            1 for item in ranking if item.Selection.status == "WAITING"
        )
        rejected_count = sum(
            1 for item in ranking if item.Selection.status == "REJECTED"
        )
        summary_rows.append([
            index,
            course.name,
            course.instructor or "",
            course.schedule or "",
            course.location or "",
            course.capacity,
            len(ranking),
            admitted_count,
            waiting_count,
            rejected_count,
            "是" if admitted_count > course.capacity else "否",
            course.description or "",
        ])
        detail_rows.extend(_ranking_rows(course, ranking))

    _write_table(
        summary,
        ["序号", "课程名称", "教师", "上课时间", "地点", "课程容量", "选课记录数", "已录取/确认", "候补人数", "未录取人数", "是否超员", "课程描述"],
        summary_rows,
        [8, 24, 14, 22, 18, 12, 12, 14, 12, 12, 12, 36],
    )
    for row in range(5, 5 + len(summary_rows)):
        if summary.cell(row=row, column=11).value == "是":
            for column in range(1, 13):
                summary.cell(row=row, column=column).fill = WARNING_FILL

    details = workbook.create_sheet("选课明细")
    _style_title(details, f"{period.name} · 学生选课明细", 11)
    _style_metadata(details, _period_metadata(period), 11)
    _write_table(
        details,
        ["课程名称", "教师", "上课时间", "地点", "排名", "学号", "姓名", "权重", "选课状态", "选课时间", "备注"],
        detail_rows,
        [24, 14, 22, 18, 9, 18, 14, 10, 12, 22, 14],
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_course_selection_workbook(period, course, ranking):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "选课名单"
    _style_title(sheet, f"{course.name} · 选课情况", 7)

    admitted_count = sum(
        1 for item in ranking if item.Selection.status in ("SELECTED", "FINAL")
    )
    waiting_count = sum(
        1 for item in ranking if item.Selection.status == "WAITING"
    )
    rejected_count = sum(
        1 for item in ranking if item.Selection.status == "REJECTED"
    )
    period_name = period.name if period else "未分配阶段"
    metadata = (
        f"选课阶段：{period_name}    课程容量：{course.capacity}    "
        f"选课记录数：{len(ranking)}    已录取/确认：{admitted_count}    "
        f"候补：{waiting_count}    未录取：{rejected_count}    "
        f"教师：{course.instructor or '-'}    时间：{course.schedule or '-'}    "
        f"地点：{course.location or '-'}"
    )
    _style_metadata(sheet, metadata, 7)

    rows = [row[4:] for row in _ranking_rows(course, ranking)]
    _write_table(
        sheet,
        ["排名", "学号", "姓名", "权重", "选课状态", "选课时间", "备注"],
        rows,
        [9, 18, 14, 10, 12, 22, 14],
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
