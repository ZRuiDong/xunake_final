from dataclasses import dataclass
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


EXPECTED_HEADERS = ("学号", "姓名")
MAX_IMPORT_ROWS = 2000


@dataclass(frozen=True)
class StudentImportRecord:
    row_number: int
    student_no: str
    name: str


class StudentImportValidationError(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("student import validation failed")
        self.errors = errors


def build_student_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "学生导入"
    sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="0F766E")
    for column, value in enumerate(EXPECTED_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    sheet["A1"].comment = Comment("必填；请将学号列设置为文本，避免丢失前导零。", "系统")
    sheet["B1"].comment = Comment("必填；最长 50 个字符。", "系统")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["A"].number_format = "@"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_student_workbook(content: bytes):
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise StudentImportValidationError(["无法读取文件，请确认文件是有效的 .xlsx 文件"]) from error

    sheet = workbook.active
    headers = tuple(normalize_text(sheet.cell(1, column).value) for column in range(1, 3))
    errors = []

    if headers != EXPECTED_HEADERS:
        errors.append("第 1 行表头必须依次为：学号、姓名")

    if any(
        normalize_text(sheet.cell(1, column).value)
        for column in range(3, sheet.max_column + 1)
    ):
        errors.append("模板只能包含学号、姓名两列")

    records = []
    seen_student_numbers = set()
    data_row_count = 0

    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        first_two = tuple(values[:2]) + (None,) * max(0, 2 - len(values))
        student_no = normalize_text(first_two[0])
        name = normalize_text(first_two[1])
        extra_values = values[2:]

        if not student_no and not name and not any(
            normalize_text(value) for value in extra_values
        ):
            continue

        data_row_count += 1
        if data_row_count > MAX_IMPORT_ROWS:
            errors.append(f"最多允许导入 {MAX_IMPORT_ROWS} 名学生")
            break

        row_errors = []
        if not student_no:
            row_errors.append("学号不能为空")
        elif len(student_no) > 50:
            row_errors.append("学号不能超过 50 个字符")
        elif student_no in seen_student_numbers:
            row_errors.append("学号在文件中重复")

        if not name:
            row_errors.append("姓名不能为空")
        elif len(name) > 50:
            row_errors.append("姓名不能超过 50 个字符")

        if any(normalize_text(value) for value in extra_values):
            row_errors.append("只能填写学号、姓名两列")

        if row_errors:
            errors.append(f"第 {row_number} 行：{'；'.join(row_errors)}")
            continue

        seen_student_numbers.add(student_no)
        records.append(StudentImportRecord(
            row_number=row_number,
            student_no=student_no,
            name=name,
        ))

    workbook.close()

    if not records and not errors:
        errors.append("文件中没有学生数据")

    if errors:
        raise StudentImportValidationError(errors)

    return records
