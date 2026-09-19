import unittest
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.admin.selection_export import (
    build_course_selection_workbook,
    build_period_selection_workbook,
)
from app.database import Base
from app.models import Course, Period, Selection, Student, User
from app.utils.selection import calculate_ranking
from app.utils.clock import app_now


class SelectionExportTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

        self.period = Period(
            name="秋季兴趣课选课",
            start_time=app_now() - timedelta(days=1),
            end_time=app_now() + timedelta(days=1),
            status="ACTIVE",
        )
        self.db.add(self.period)
        self.db.flush()

        self.course = Course(
            name="绘画基础",
            description="水彩与素描入门",
            capacity=1,
            status="OPEN",
            period_id=self.period.id,
        )
        self.empty_course = Course(
            name="篮球基础",
            capacity=20,
            status="OPEN",
            period_id=self.period.id,
        )
        self.db.add_all([self.course, self.empty_course])
        self.db.flush()

        statuses = ["FINAL", "WAITING", "REJECTED"]
        for index, status in enumerate(statuses, start=1):
            user = User(
                username=f"S{index:03d}",
                password_hash="test",
                role="STUDENT",
            )
            self.db.add(user)
            self.db.flush()
            student = Student(
                user_id=user.id,
                student_no=f"S{index:03d}",
                name=f"测试学生{index}",
            )
            self.db.add(student)
            self.db.flush()
            self.db.add(Selection(
                student_id=student.id,
                course_id=self.course.id,
                period_id=self.period.id,
                selected_time=datetime(2026, 8, 26, 9, index),
                status=status,
            ))
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_period_workbook_contains_summary_and_all_selection_details(self):
        courses = [self.course, self.empty_course]
        rankings = {
            course.id: calculate_ranking(course.id, self.db, include_rejected=True)
            for course in courses
        }
        content = build_period_selection_workbook(self.period, courses, rankings)
        workbook = load_workbook(BytesIO(content), data_only=True)

        self.assertEqual(workbook.sheetnames, ["课程汇总", "选课明细"])
        summary = workbook["课程汇总"]
        details = workbook["选课明细"]
        self.assertEqual(summary["A1"].value, "秋季兴趣课选课 · 选课情况汇总")
        self.assertEqual(summary["B5"].value, "绘画基础")
        self.assertEqual(summary["G5"].value, 3)
        self.assertEqual(summary["J5"].value, 1)
        self.assertEqual(summary["B6"].value, "篮球基础")
        self.assertEqual(details["F5"].value, "S001")
        self.assertEqual(details["H6"].value, "候补中")
        self.assertEqual(details["H7"].value, "未录取")
        workbook.close()

    def test_course_workbook_contains_course_overview_and_ranked_students(self):
        ranking = calculate_ranking(
            self.course.id,
            self.db,
            include_rejected=True,
        )
        content = build_course_selection_workbook(
            self.period,
            self.course,
            ranking,
        )
        workbook = load_workbook(BytesIO(content), data_only=True)
        sheet = workbook["选课名单"]

        self.assertEqual(sheet["A1"].value, "绘画基础 · 选课情况")
        self.assertIn("课程容量：1", sheet["A2"].value)
        self.assertEqual(sheet["B5"].value, "S001")
        self.assertEqual(sheet["D6"].value, "候补中")
        self.assertEqual(sheet["F6"].value, "候补")
        self.assertEqual(sheet["A7"].value, "-")
        self.assertEqual(sheet["D7"].value, "未录取")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
