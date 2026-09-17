import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.admin.selection_export import _write_table, _style_title


class ExportSafetyTest(unittest.TestCase):
    def test_user_controlled_strings_are_not_excel_formulas(self):
        workbook = Workbook()
        sheet = workbook.active
        _style_title(sheet, "=1+1", 2)
        _write_table(sheet, ["name", "student no"], [["=1+1", "00123"]], [20, 20])
        content = BytesIO()
        workbook.save(content)
        content.seek(0)
        sheet = load_workbook(content).active
        for address in ("A1", "A5", "B5"):
            self.assertEqual(sheet[address].data_type, "s")
        self.assertEqual(sheet["A5"].value, "=1+1")
        self.assertEqual(sheet["B5"].value, "00123")
