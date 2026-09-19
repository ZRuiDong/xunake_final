import unittest
from io import BytesIO

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.admin.router import import_students
from app.admin.student_import import (
    StudentImportValidationError,
    build_student_template,
    parse_student_workbook,
)
from app.auth.security import verify_password
from app.database import Base
from app.models import Student, User


def workbook_bytes(rows, headers=("学号", "姓名")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class StudentImportTest(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = sessionmaker(bind=self.engine)()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_download_template_has_exactly_two_headers(self):
        workbook = load_workbook(BytesIO(build_student_template()), read_only=True)
        sheet = workbook.active
        headers = tuple(sheet.cell(1, column).value for column in range(1, 3))
        self.assertEqual(headers, ("学号", "姓名"))
        self.assertEqual(sheet.max_column, 2)
        workbook.close()

    def test_parser_reports_row_number_for_duplicate_student_number(self):
        content = workbook_bytes([
            ("2027001", "学生甲"),
            ("2027001", "学生乙"),
        ])

        with self.assertRaises(StudentImportValidationError) as context:
            parse_student_workbook(content)

        self.assertIn("第 3 行", context.exception.errors[0])
        self.assertIn("重复", context.exception.errors[0])

    def test_old_three_column_template_is_rejected(self):
        content = workbook_bytes(
            [("2027001", "学生甲", 10)],
            headers=("学号", "姓名", "权重"),
        )

        with self.assertRaises(StudentImportValidationError) as context:
            parse_student_workbook(content)

        self.assertTrue(any("两列" in error for error in context.exception.errors))

    async def test_valid_workbook_imports_all_students(self):
        content = workbook_bytes([
            ("2027001", "学生甲"),
            ("2027002", "学生乙"),
        ])
        upload = UploadFile(filename="students.xlsx", file=BytesIO(content))

        result = await import_students(file=upload, db=self.db, admin="admin")

        self.assertEqual(result["imported_count"], 2)
        self.assertEqual(self.db.query(Student).count(), 2)
        user = self.db.query(User).filter(User.username == "2027001").one()
        self.assertTrue(verify_password("123456", user.password_hash))
        self.assertTrue(user.must_change_password)

    async def test_existing_student_rejects_whole_file(self):
        user = User(username="2027001", password_hash="test", role="STUDENT")
        self.db.add(user)
        self.db.flush()
        self.db.add(Student(
            user_id=user.id,
            student_no="2027001",
            name="已存在学生",
        ))
        self.db.commit()

        content = workbook_bytes([
            ("2027001", "重复学生"),
            ("2027002", "新学生"),
        ])
        upload = UploadFile(filename="students.xlsx", file=BytesIO(content))

        with self.assertRaises(HTTPException) as context:
            await import_students(file=upload, db=self.db, admin="admin")

        self.assertEqual(context.exception.status_code, 400)
        self.assertEqual(self.db.query(Student).count(), 1)


if __name__ == "__main__":
    unittest.main()
